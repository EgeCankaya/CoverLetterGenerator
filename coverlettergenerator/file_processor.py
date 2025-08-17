"""File processing utilities for CV uploads."""

from typing import ClassVar, Optional

from docx import Document
from PyPDF2 import PdfReader


class FileProcessor:
    """Handles file uploads and text extraction from CV files."""

    ALLOWED_EXTENSIONS: ClassVar[set[str]] = {".pdf", ".docx"}

    def is_valid_file(self, filename: str) -> bool:
        """Check if the uploaded file has a valid extension."""
        if not filename:
            return False

        file_extension = filename.lower().rsplit(".", 1)[1] if "." in filename else ""
        return f".{file_extension}" in self.ALLOWED_EXTENSIONS

    def extract_text(self, file) -> Optional[str]:
        """Extract text content from uploaded CV file."""
        try:
            filename = file.filename.lower()

            if filename.endswith(".pdf"):
                return self._extract_from_pdf(file)
            elif filename.endswith(".docx"):
                return self._extract_from_docx(file)
            else:
                return None

        except Exception as e:
            print(f"Error extracting text from file: {e!s}")
            return None

    def _extract_from_pdf(self, file) -> Optional[str]:
        """Extract text from PDF file."""
        try:
            # Read the PDF file
            pdf_reader = PdfReader(file)
            text_content = []

            # Extract text from each page
            for page in pdf_reader.pages:
                text_content.append(page.extract_text())

            return "\n".join(text_content)

        except Exception as e:
            print(f"Error extracting text from PDF: {e!s}")
            return None

    def _extract_from_docx(self, file) -> Optional[str]:
        """Extract text from DOCX file."""
        try:
            # Read the DOCX file
            doc = Document(file)
            text_content = []

            # Extract text from each paragraph
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content.append(paragraph.text)

            # Also extract text from tables if present
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        text_content.append(" | ".join(row_text))

            return "\n".join(text_content)

        except Exception as e:
            print(f"Error extracting text from DOCX: {e!s}")
            return None
